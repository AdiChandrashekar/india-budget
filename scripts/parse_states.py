"""Turn RBI's state finance workbooks into the project's long table.

    python scripts/parse_states.py

Writes

    data/processed/state_education.csv     the long table, same shape as the Union's
    data/processed/state_statement26.csv   RBI's published percentages, as fixtures
    data/reference/state_gsdp.json         GSDP, for ratios

The source is the e-STATES Database that ships with *State Finances: A Study
of Budgets*. It is a single tidy sheet of 397,575 rows -- appendix, state,
budget head, fiscal year, then three money columns for accounts, revised and
budget estimates -- covering 1990-91 to 2025-26 for 31 states and union
territories with legislatures, plus RBI's own all-states aggregate.

That one file carries the whole span, so there is nothing to stitch: the plan
to splice several editions together for depth turned out to be unnecessary.
See docs/FINDINGS.md.

What it does *not* carry is sub-major-head detail. RBI reports the functional
group "Education, Sports, Art and Culture", not elementary against secondary
against higher education. On the expenditure side that group is major heads
2202, 2203, 2204 and 2205 together; capital outlay (4202), loans (6202) and
receipts (0202) are single major heads, so those three are exact. The group is
the same functional definition the Union view already uses, which is the point
-- it is what makes the two halves comparable.
"""

from __future__ import annotations

import collections
import re
import sys
import time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

import common as c  # noqa: E402

RAW = c.RAW / "states"
EDITION_LABEL = "RBI State Finances %s (e-STATES)"

# --------------------------------------------------------------------------- #
# What to read out of the e-STATES sheet
# --------------------------------------------------------------------------- #
# (appendix, budget head) -> (definition, account, major_head, label)
#
# The appendices are the publication's four statistical appendices: revenue
# receipts, revenue expenditure, capital receipts, capital disbursements. The
# same function therefore appears under several of them meaning quite different
# things, and the one under Appendix-1 is money coming *in*. Keying on the head
# label alone would add fee receipts to spending.
HEADS = {
    # ---- education, the functional group -----------------------------------
    ("Appendix-2", "I.A.1: Education, Sports, Art and Culture"):
        ("functional", "revenue_expenditure", "",
         "Education, Sports, Art and Culture"),
    ("Appendix-4", "I.1.a.1: Education, Sports, Art and Culture"):
        ("functional", "capital_outlay", "4202",
         "Capital Outlay on Education, Sports, Art and Culture"),
    ("Appendix-4", "IV.1.a.1: Education, Sports, Art and Culture"):
        ("functional", "loans_and_advances", "6202",
         "Loans for Education, Sports, Art and Culture"),
    # Receipts, not spending. Carried so the netting can be shown rather than
    # guessed at, and never added to the expenditure lines.
    ("Appendix-1", "II.C.4.i: Education, Sports, Art and Culture"):
        ("functional", "revenue_receipts", "0202",
         "Education, Sports, Art and Culture (receipts)"),

    # ---- the denominators, and the social sector around education ----------
    ("Appendix-2", "Total: TOTAL EXPENDITURE (I+II+III)"):
        ("whole", "revenue_expenditure", "", "Total revenue expenditure"),
    ("Appendix-2", "I.A: Social Services (1 to 12)"):
        ("whole", "revenue_expenditure", "", "Social Services"),
    ("Appendix-4", "I: Total Capital Outlay (1 + 2)"):
        ("whole", "capital_outlay", "", "Total capital outlay"),
    ("Appendix-4", "IV: Loans and Advances by State Governments (1+2)"):
        ("whole", "loans_and_advances", "", "Total loans and advances"),
    # RBI's "aggregate expenditure", the denominator of Statement 26, is this
    # line plus total revenue expenditure. Not capital outlay plus loans, and
    # not the "TOTAL CAPITAL DISBURSEMENTS (I to XII)" line either -- that one
    # runs to eleven million crore because it sweeps in the public account.
    # Reproducing all 570 published percentages is what settled which it was.
    ("Appendix-4",
     "Total$: TOTAL CAPITAL DISBURSEMENTS (Excluding Public Accounts)"):
        ("whole", "capital_disbursement", "",
         "Total capital disbursements, excluding the public account"),
}

# The three money columns, left to right.
ESTIMATE_COLUMNS = [(4, "A"), (5, "RE"), (6, "BE")]

# The head whose absence means the entity had no budget that year. See below.
SCOPE_HEAD = ("Appendix-2", "Total: TOTAL EXPENDITURE (I+II+III)")


def money(value):
    """A money cell, or None. e-STATES stores most numbers as text."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip()
    if s in ("", "-", "--", "–", "—"):
        return None
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def load_states():
    ref = c.read_json(c.REFERENCE / "states.json")
    by_rbi = {s["rbi_name"]: s for s in ref["states"]}
    return ref, by_rbi


# --------------------------------------------------------------------------- #
# e-STATES
# --------------------------------------------------------------------------- #


def read_estates(path: Path):
    """[(appendix, rbi_state, head, fiscal_year, [(etype, amount)])]"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Data" not in wb.sheetnames:
        raise SystemExit("%s has no Data sheet; sheets are %s"
                         % (path.name, wb.sheetnames))
    rows = wb["Data"].iter_rows(values_only=True)
    header = [c.squeeze(v) for v in next(rows)]
    expected = ["Appendix", "State/UT", "Budget Head", "Fiscal Year",
                "Account", "Revised", "Budget"]
    if header[:7] != expected:
        raise SystemExit("e-STATES columns have moved: expected %s, got %s"
                         % (expected, header[:7]))

    out = []
    for row in rows:
        if row[0] is None:
            continue
        key = (c.squeeze(row[0]), c.squeeze(row[2]))
        if key not in HEADS:
            continue
        amounts = [(etype, money(row[i])) for i, etype in ESTIMATE_COLUMNS]
        out.append((key[0], c.squeeze(row[1]), key[1],
                    c.fy_normalise(row[3]), amounts))
    if not out:
        raise SystemExit("none of the wanted heads appear in %s" % path.name)
    return out


def in_scope(records):
    """Which (state, year, estimate type) the entity actually had a budget for.

    This is the trap that makes a state layer lie quietly. e-STATES pads a
    state's pre-formation years with a literal 0 rather than leaving them
    blank, for every head, in every appendix. Parse it naively and Chhattisgarh
    spends Rs 0 crore on education in 1995 -- six years before it existed -- a
    chart draws a confident line along the axis to 2000, and every average over
    "all states" is quietly diluted by states that were not there.

    A zero cannot simply be dropped, because a genuine nil is possible on a
    small head: several states really do lend nothing for education in a given
    year. So the test is applied to the entity, not the number. If a state's
    total revenue expenditure is zero, it had no budget on that basis and the
    whole cell goes; where it did have one, a zero on a particular head is a
    real nil and is kept.

    The same padding also appears per estimate type rather than per year.
    Puducherry's 2005-06 accounts are published, but its revised and budget
    estimates for that year are zeros, so the test has to be applied to each
    of the three money columns separately -- keying it on the year alone
    readmits exactly the empty cells it was meant to exclude.
    """
    scope = set()
    for appendix, state, head, fy, amounts in records:
        if (appendix, head) != SCOPE_HEAD:
            continue
        for etype, amount in amounts:
            if amount is not None and amount > 0:
                scope.add((state, fy, etype))
    return scope


def check_formation(scope, ref, by_rbi) -> list:
    """Cross-check the reference table's formation years against the data."""
    first = {}
    for state, fy, _etype in scope:
        if state not in first or c.fy_start(fy) < c.fy_start(first[state]):
            first[state] = fy
    problems = []
    for rbi_name, fy in sorted(first.items()):
        meta = by_rbi.get(rbi_name)
        if meta is None:
            problems.append("e-STATES has %r, which states.json does not"
                            % rbi_name)
            continue
        if c.fy_normalise(meta["formed"]) != fy:
            problems.append(
                "%s: states.json says formed %s, the data starts at %s"
                % (meta["code"], meta["formed"], fy))
    for s in ref["states"]:
        if s["rbi_name"] not in first:
            problems.append("states.json has %s, which e-STATES does not"
                            % s["code"])
    return problems


def to_rows(records, scope, by_rbi, edition: str, url: str):
    rows = []
    occurrence = collections.Counter()
    for appendix, state, head, fy, amounts in records:
        meta = by_rbi.get(state)
        if meta is None:
            continue
        definition, account, major_head, label = HEADS[(appendix, head)]
        for etype, amount in amounts:
            if amount is None or (state, fy, etype) not in scope:
                continue
            key = (fy, meta["code"], etype, definition, account, major_head,
                   label)
            occurrence[key] += 1
            rows.append(c.Row(
                fiscal_year=fy,
                govt="state",
                state_code=meta["code"],
                estimate_type=etype,
                source_edition=edition,
                definition=definition,
                account=account,
                major_head=major_head,
                scheme=label,
                amount_cr=c.to_crore(amount, "crore"),
                source_url=url,
                source_page="estates!%s!%s" % (appendix, head.split(":")[0]),
                occurrence=occurrence[key],
            ))
    return rows


# --------------------------------------------------------------------------- #
# Statement 26, kept as fixtures rather than as data
# --------------------------------------------------------------------------- #

# The final column is headed "2025-26 (RE)" in the published workbook. There is
# no revised estimate for 2025-26 anywhere in e-STATES, the volume is the study
# of the 2025-26 *budgets*, and the e-STATES notes say the database runs "to
# 2025-26 (Budget Estimates)". It is a typo for (BE), and reading it as RE
# leaves the whole last column unmatchable. See docs/FINDINGS.md.
LAST_COLUMN_IS_BE = "2025-26"


def read_statement26(path: Path, by_rbi):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    header, header_at = None, None
    for i, row in enumerate(rows[:12]):
        cells = [c.squeeze(v) for v in row if v is not None]
        if cells and cells[0] in ("State/UT", "State / UT"):
            header, header_at = cells[1:], i
            break
    if header is None:
        raise SystemExit("could not find the header row in %s" % path.name)

    out = []
    for row in rows[header_at + 1:]:
        cells = [v for v in row if v is not None]
        if not cells:
            continue
        name = re.sub(r"^\d+\.\s*", "", c.squeeze(cells[0]))
        if name == "All States and UTs":
            name = "All States/UT"
        meta = by_rbi.get(name)
        if meta is None:
            continue
        for j, label in enumerate(header):
            value = money(cells[1 + j]) if 1 + j < len(cells) else None
            if value is None:
                continue
            fy = c.fy_normalise(label[:7])
            if label.startswith(LAST_COLUMN_IS_BE):
                etype = "BE"
            elif "(RE)" in label:
                etype = "RE"
            elif "(BE)" in label:
                etype = "BE"
            else:
                etype = "A"
            out.append({
                "state_code": meta["code"],
                "fiscal_year": fy,
                "estimate_type": etype,
                "published_pct": value,
                "column_header": label,
            })
    if not out:
        raise SystemExit("no percentages read from %s" % path.name)
    return out


# --------------------------------------------------------------------------- #
# GSDP
# --------------------------------------------------------------------------- #

# The Handbook spells some states differently from the state finances volume,
# and carries three entities the budget series has no row for.
GSDP_NAMES = {
    "Jammu & Kashmir*": "Jammu and Kashmir",
    "Delhi": "NCT Delhi",
}
# No budget of their own in RBI's state finances series, so no ratio to form.
GSDP_SKIP = {"Andaman & Nicobar Islands", "Chandigarh", "Ladakh"}


def read_gsdp(path: Path, by_rbi):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    series = collections.defaultdict(dict)
    unit, notes = "", []
    for name in wb.sheetnames:
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
        years, years_at = None, None
        for i, row in enumerate(rows[:10]):
            cells = [c.squeeze(v) for v in row if v is not None]
            if len(cells) == 1 and cells[0].startswith("("):
                unit = unit or cells[0].strip("()")
            if cells and re.fullmatch(r"(?:19|20)\d{2}-\d{2}", cells[0]):
                years, years_at = cells, i
                break
        if years is None:
            continue
        for row in rows[years_at + 1:]:
            cells = [v for v in row if v is not None]
            if len(cells) < 2:
                if cells:
                    text = c.squeeze(cells[0])
                    if text.startswith(("-:", "*:", "Source")):
                        notes.append(text)
                continue
            label = c.squeeze(cells[0])
            if label in GSDP_SKIP:
                continue
            meta = by_rbi.get(GSDP_NAMES.get(label, label))
            if meta is None:
                continue
            for j, year in enumerate(years):
                value = money(cells[1 + j]) if 1 + j < len(cells) else None
                if value is None:
                    continue
                # The Handbook prints GSDP in lakh; everything here is crore.
                series[meta["code"]][c.fy_normalise(year)] = round(
                    c.to_crore(value, "lakh"), 2)
    return series, unit, sorted(set(notes))


# --------------------------------------------------------------------------- #


def main() -> None:
    ref, by_rbi = load_states()
    editions = c.read_json(c.REFERENCE / "state_editions.json")
    docs = dict(editions["state_finances"]["documents"])
    docs.update(editions["handbook"]["documents"])
    edition = EDITION_LABEL % (editions.get("edition") or "unknown")

    estates_path = RAW / "estates.xlsx"
    if not estates_path.exists():
        raise SystemExit("no %s; run scripts/fetch_states.py first"
                         % estates_path.name)

    print("reading %s" % estates_path.name)
    records = read_estates(estates_path)
    print("  %s rows on the heads we want" % format(len(records), ","))

    scope = in_scope(records)
    print("  %d state-years in scope" % len(scope))

    problems = check_formation(scope, ref, by_rbi)
    if problems:
        print("\nthe reference table and the data disagree:")
        for p in problems:
            print("  -", p)
        raise SystemExit(
            "refusing to write: states.json records a formation year the data "
            "does not support, so either the reference table is wrong or RBI "
            "has changed what it publishes")
    print("  formation years agree with data/reference/states.json")

    rows = to_rows(records, scope, by_rbi,
                   edition, docs["estates"]["url"])
    c.write_long_table(rows, c.PROCESSED / "state_education.csv")

    # ---- the published percentages, as a committed fixture set -------------
    st26 = RAW / "statement_26.xlsx"
    if st26.exists():
        fixtures = read_statement26(st26, by_rbi)
        dest = c.PROCESSED / "state_statement26.csv"
        import csv as _csv
        tmp = dest.with_suffix(".csv.tmp")
        cols = ["state_code", "fiscal_year", "estimate_type", "published_pct",
                "column_header", "source_url"]
        url = docs["statement_26"]["url"]
        with tmp.open("w", newline="", encoding="utf-8") as fh:
            w = _csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            for f in fixtures:
                f = dict(f, source_url=url)
                w.writerow({k: f[k] for k in cols})
        tmp.replace(dest)
        print("wrote %s published percentages -> %s"
              % (format(len(fixtures), ","), dest.relative_to(c.ROOT)))
    else:
        print("! no statement_26.xlsx; the published fixtures were not refreshed")

    # ---- GSDP --------------------------------------------------------------
    gsdp_path = RAW / "handbook_t21_gsdp_current.xlsx"
    if gsdp_path.exists():
        series, unit, notes = read_gsdp(gsdp_path, by_rbi)
        years = sorted({y for s in series.values() for y in s}, key=c.fy_start)
        c.write_json({
            "generated": time.strftime("%Y-%m-%d", time.gmtime()),
            "source": docs["gsdp_current"]["label"],
            "source_url": docs["gsdp_current"]["url"],
            "published_unit": unit,
            "unit": "Rs crore, current prices, base 2011-12",
            "years": years,
            "notes": notes,
            "series": {k: series[k] for k in sorted(series)},
        }, c.REFERENCE / "state_gsdp.json", indent=0)
        print("GSDP: %d entities, %s to %s"
              % (len(series), years[0], years[-1]))
    else:
        print("! no GSDP table; state_gsdp.json was not refreshed")

    # ---- what we ended up with --------------------------------------------
    years = sorted({r.fiscal_year for r in rows}, key=c.fy_start)
    codes = sorted({r.state_code for r in rows})
    print("\n%s rows, %d entities, %s to %s"
          % (format(len(rows), ","), len(codes), years[0], years[-1]))
    by_type = collections.Counter(r.estimate_type for r in rows)
    print("estimate types: %s"
          % ", ".join("%s %s" % (k, format(v, ",")) for k, v in
                      sorted(by_type.items())))


if __name__ == "__main__":
    main()
